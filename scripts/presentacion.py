#!/usr/bin/env python3
"""Genera presentacion/preguntazo.pptx con capturas reales del producto.

Requiere python-pptx y Pillow (pip install python-pptx pillow). Las imágenes salen de
presentacion/img, que llena scripts/capture-store.mjs con capturas reales de la
extensión, el celular, el panel y la presentación de prueba con la capa encima.

    python3 scripts/presentacion.py
"""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "presentacion" / "img"
OUT = ROOT / "presentacion" / "preguntazo.pptx"
GEN = IMG / "generadas"
GEN.mkdir(parents=True, exist_ok=True)

BG = RGBColor(0x0B, 0x0B, 0x0C)
CARD = RGBColor(0x18, 0x18, 0x1B)
CARD2 = RGBColor(0x22, 0x22, 0x26)
TEXT = RGBColor(0xF5, 0xF5, 0xF7)
MUTED = RGBColor(0xA1, 0xA1, 0xA6)
ACCENT = RGBColor(0xBF, 0x5A, 0xF2)
FONT = "Arial"

# ---------------------------------------------------------------- imágenes

def fondo():
    path = GEN / "fondo.png"
    w, h = 1920, 1080
    base = Image.new("RGB", (w, h), (11, 11, 12))
    glow = Image.new("L", (w, h), 0)
    d = ImageDraw.Draw(glow)
    d.ellipse((1150, 520, 2250, 1400), fill=120)
    d.ellipse((-400, -500, 600, 300), fill=40)
    glow = glow.filter(ImageFilter.GaussianBlur(220))
    cyan = Image.new("RGB", (w, h), (191, 90, 242))
    base = Image.composite(cyan, base, glow.point(lambda v: int(v * 0.3)))
    base.save(path)
    return path


def redondear(src, dest, radio):
    im = Image.open(src).convert("RGBA")
    mask = Image.new("L", im.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, im.width - 1, im.height - 1), radius=radio, fill=255)
    im.putalpha(mask)
    im.save(dest)
    return dest


def celular(nombre):
    """Pantalla del celular con bordes redondeados y marco oscuro."""
    dest = GEN / f"{nombre}-marco.png"
    pantalla = Image.open(IMG / f"{nombre}.png").convert("RGBA")
    pantalla = pantalla.resize((pantalla.width * 1170 // pantalla.width, pantalla.height * 1170 // pantalla.width), Image.LANCZOS)
    sw, sh = pantalla.size
    marco = 34
    w, h = sw + marco * 2, sh + marco * 2
    out = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    cuerpo = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(cuerpo)
    d.rounded_rectangle((0, 0, w - 1, h - 1), radius=150, fill=(26, 26, 29, 255), outline=(70, 70, 76, 255), width=3)
    out.alpha_composite(cuerpo)
    mask = Image.new("L", (sw, sh), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, sw - 1, sh - 1), radius=118, fill=255)
    pantalla.putalpha(mask)
    out.alpha_composite(pantalla, (marco, marco))
    out.save(dest)
    return dest


# ---------------------------------------------------------------- utilidades de diapositivas

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
FONDO = fondo()


def nueva():
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(str(FONDO), 0, 0, prs.slide_width, prs.slide_height)
    return s


def texto(s, x, y, w, h, contenido, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line=1.15):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    parrafos = contenido if isinstance(contenido, list) else [contenido]
    for i, p_text in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if i:
            p.space_before = Pt(size * 0.5)
        run = p.add_run()
        run.text = p_text
        f = run.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        if spacing is not None:
            rPr = run._r.get_or_add_rPr()
            rPr.set("spc", str(spacing))
    return box


def cabecera(s, kicker, titulo, y=0.62, w=11.8, size=36):
    texto(s, 0.7, y, 8, 0.3, kicker.upper(), size=12, color=ACCENT, bold=True, spacing=200)
    texto(s, 0.7, y + 0.38, w, 1.2, titulo, size=size, bold=True, line=1.05)


def tarjeta(s, x, y, w, h, titulo, cuerpo, numero=None, color=CARD):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    top = y + 0.3
    if numero is not None:
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.3), Inches(top), Inches(0.42), Inches(0.42))
        c.fill.solid()
        c.fill.fore_color.rgb = ACCENT
        c.line.fill.background()
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(numero)
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        top += 0.62
    # Arial negrita de 17 pt: unos 6 caracteres por pulgada. Si el título ocupa dos
    # líneas, el texto empieza más abajo.
    lineas = 2 if len(titulo) > (w - 0.6) * 6.2 else 1
    texto(s, x + 0.3, top, w - 0.6, 0.3 * lineas + 0.2, titulo, size=17, bold=True, line=1.05)
    cuerpo_y = top + 0.18 + 0.3 * lineas
    texto(s, x + 0.3, cuerpo_y, w - 0.6, h - (cuerpo_y - y) - 0.2, cuerpo, size=13, color=MUTED, line=1.25)
    return shp


def foto(s, path, x, y, w=None, h=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def lista(s, x, y, w, items, size=15, gap=0.62):
    for i, (titulo, detalle) in enumerate(items):
        yy = y + i * gap
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy + 0.09), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT
        dot.line.fill.background()
        texto(s, x + 0.3, yy, w - 0.3, 0.3, titulo, size=size, bold=True)
        texto(s, x + 0.3, yy + 0.27, w - 0.3, 0.3, detalle, size=size - 3, color=MUTED)


# ---------------------------------------------------------------- imágenes preparadas

phones = {n: celular(n) for n in ["phone-bienvenida", "phone-preguntas", "phone-enviada", "phone-encuesta", "phone-resultados", "phone-nube", "phone-reacciones"]}
r = lambda n, radio=36: redondear(IMG / f"{n}.png", GEN / f"{n}-r.png", radio)
s_reacciones = r("slides-reacciones")
s_pregunta = r("slides-pregunta")
s_encuesta = r("slides-encuesta")
s_nube = r("slides-nube")
s_qr = r("slides-qr-grande")
s_dos = r("slides-diapositiva-2")
pantalla = r("pantalla")
panel_q = r("panel-preguntas", 28)
panel_p = r("panel-encuestas", 28)
panel_r = r("panel-reacciones", 28)
popup = r("popup", 34)
bienvenida = r("options", 28)

# ---------------------------------------------------------------- 1. portada

s = nueva()
foto(s, IMG / "logo-512.png", 0.7, 1.55, w=1.25)
texto(s, 0.7, 3.0, 7.4, 1.1, "Preguntazo", size=66, bold=True)
texto(s, 0.7, 4.15, 7.0, 1.2, "Preguntas, encuestas y reacciones en vivo para tus charlas", size=24, color=MUTED, line=1.2)
texto(s, 0.7, 6.55, 7, 0.4, "Miguel Mogrovejo · Octubre de 2026", size=13, color=MUTED)
foto(s, phones["phone-preguntas"], 9.15, 0.55, h=6.4)

# ---------------------------------------------------------------- 2. el reto

s = nueva()
cabecera(s, "El reto", "El público escucha, pero casi nunca participa")
for i, (t, c) in enumerate([
    ("Nadie levanta la mano", "Da vergüenza preguntar frente a todos, y las mejores preguntas se quedan sin hacer."),
    ("Las preguntas se pierden", "El micrófono no llega, el tiempo se acaba y siempre preguntan los mismos."),
    ("Hablas a ciegas", "No sabes si te siguen, qué piensan o qué les interesa hasta el final."),
    ("Las apps estorban", "Descargar algo, registrarse o crear una cuenta hace que nadie participe."),
]):
    tarjeta(s, 0.7 + i * 3.03, 2.25, 2.83, 3.6, t, c, numero=i + 1)

# ---------------------------------------------------------------- 3. la solución

s = nueva()
cabecera(s, "La solución", "Un QR, y toda la sala participa", w=6.0)
texto(s, 0.7, 2.55, 4.9, 1.3, "Una extensión de Chrome y una página web: tu público escanea el QR de tus diapositivas y participa desde el navegador de su celular.", size=16, color=MUTED, line=1.3)
lista(s, 0.7, 3.95, 4.9, [
    ("Preguntas con votos", "Las más votadas suben solas y las muestras con un clic."),
    ("Encuestas en vivo", "Una opción, varias, estrellas y nube de palabras."),
    ("Reacciones flotantes", "Emojis que suben sobre tu presentación en tiempo real."),
], gap=0.78)
foto(s, s_reacciones, 6.45, 1.5, w=6.35)
foto(s, phones["phone-reacciones"], 6.05, 3.1, h=4.1)

# ---------------------------------------------------------------- 4. cómo se usa

s = nueva()
cabecera(s, "Cómo se usa", "Tres pasos, sin apps ni cuentas")
for i, (t, c) in enumerate([
    ("Abre tu presentación", "En Chrome, haz clic en Preguntazo (o Alt+Shift+P) y activa «Mostrar en esta pestaña»."),
    ("Tu público escanea", "Un QR queda en la esquina; con Alt+Shift+Q lo muestras en grande. También pueden escribir el código."),
    ("Preguntan y votan", "Tú moderas desde el panel, destacas preguntas y abres encuestas cuando quieras."),
]):
    tarjeta(s, 0.7 + i * 2.6, 2.25, 2.45, 4.1, t, c, numero=i + 1)
foto(s, popup, 8.95, 0.55, h=6.5)

# ---------------------------------------------------------------- 5. preguntas

s = nueva()
cabecera(s, "Preguntas", "Las más votadas, primero", w=5.2)
lista(s, 0.7, 2.35, 4.9, [
    ("Con nombre o anónimas", "Cada persona elige; el nombre es opcional."),
    ("El público vota", "Las preguntas suben por votos; nadie ve quién votó."),
    ("En grande sobre la diapositiva", "«Mostrar en pantalla» o Alt+Shift+N para la siguiente."),
    ("Respuesta en el celular", "Quien preguntó recibe un aviso cuando la respondes."),
], gap=0.85)
foto(s, s_pregunta, 5.75, 1.7, w=5.0)
foto(s, phones["phone-preguntas"], 10.45, 0.6, h=6.3)

# ---------------------------------------------------------------- 6. encuestas

s = nueva()
cabecera(s, "Encuestas", "Vota desde el celular, resultados en vivo", w=6.0)
lista(s, 0.7, 2.35, 5.2, [
    ("Cuatro tipos", "Una opción, varias opciones, estrellas y nube de palabras."),
    ("Preparadas de antemano", "Créalas antes y ábrelas con un clic durante la charla."),
    ("Los celulares cambian solos", "Al abrir una encuesta, el público la ve al instante."),
    ("Resultados en pantalla", "Barras que crecen con cada voto, sobre tu diapositiva."),
], gap=0.85)
foto(s, s_encuesta, 5.75, 1.7, w=5.0)
foto(s, phones["phone-resultados"], 10.45, 0.6, h=6.3)

# ---------------------------------------------------------------- 7. nube de palabras

s = nueva()
cabecera(s, "Nube de palabras", "Lo que piensa la sala, de un vistazo", w=5.6)
texto(s, 0.7, 2.55, 4.9, 1.6, "Cada persona responde con una palabra o una frase corta. Las respuestas iguales se juntan (sin importar mayúsculas ni tildes) y las más repetidas se ven más grandes.", size=16, color=MUTED, line=1.35)
foto(s, s_nube, 0.7, 4.15, w=5.0)
foto(s, phones["phone-nube"], 6.35, 0.6, h=6.3)
foto(s, phones["phone-encuesta"], 9.5, 0.6, h=6.3)

# ---------------------------------------------------------------- 8. reacciones

s = nueva()
cabecera(s, "Reacciones", "Aplausos que se ven", w=4.6)
lista(s, 0.7, 2.35, 4.2, [
    ("Seis emojis a elección", "Aplausos, corazón, risa, fuego, asombro, pulgar arriba… o los que elijas."),
    ("Encima de todo", "Funcionan sobre la presentación en pantalla completa."),
    ("Detrás del texto", "Pasan por detrás de la pregunta o encuesta destacada."),
    ("Con límite", "Hasta 6 por segundo por celular: nadie tapa la pantalla."),
], gap=0.9)
foto(s, s_dos, 5.2, 1.6, w=7.6)

# ---------------------------------------------------------------- 9. panel del presentador

s = nueva()
cabecera(s, "Panel del presentador", "Toda la sala en una pestaña", w=7.2)
foto(s, panel_q, 0.7, 2.2, w=7.4)
lista(s, 8.45, 2.3, 4.3, [
    ("Moderación", "Revisa antes de publicar, oculta o elimina."),
    ("Fijar y responder", "Marca respondidas y fija las importantes."),
    ("Vista previa en vivo", "Lo que ve la sala, en miniatura."),
    ("Exportar CSV", "Preguntas y encuestas para después."),
    ("Público", "Quién está conectado; desconecta a quien moleste."),
], gap=0.85)

# ---------------------------------------------------------------- 10. QR y pantalla completa

s = nueva()
cabecera(s, "Para invitar a participar", "El QR, en grande cuando lo necesitas", w=8.5)
foto(s, s_qr, 0.7, 2.25, w=5.85)
foto(s, pantalla, 6.95, 2.25, w=5.85)
texto(s, 0.7, 5.75, 5.85, 1.0, ["QR grande sobre la diapositiva", "Alt+Shift+Q, con el código de 6 letras para quien no pueda escanear."], size=14, color=MUTED, line=1.25)
texto(s, 6.95, 5.75, 5.85, 1.0, ["Pantalla completa para el proyector", "Si presentas desde otro programa, o para capturar en OBS."], size=14, color=MUTED, line=1.25)

# ---------------------------------------------------------------- 11. cómo funciona

s = nueva()
cabecera(s, "Cómo funciona", "Directo del celular a tu computadora")


def nodo(x, y, w, h, titulo, detalle, color=CARD):
    tarjeta(s, x, y, w, h, titulo, detalle, color=color)


nodo(0.7, 2.5, 3.4, 2.2, "Celulares del público", "Página web: envían preguntas, votos y reacciones por WebRTC, sin instalar nada.")
nodo(4.95, 2.5, 3.4, 2.2, "Extensión", "Recibe a toda la sala, valida y limita cada mensaje, y guarda las preguntas en tu Chrome.")
nodo(9.2, 2.5, 3.4, 2.2, "Tus diapositivas", "Una capa transparente dibuja el QR, las reacciones y lo que destacas.")
for x in (4.15, 8.4):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(3.35), Inches(0.75), Inches(0.45))
    a.fill.solid()
    a.fill.fore_color.rgb = ACCENT
    a.line.fill.background()
texto(s, 0.7, 5.0, 11.9, 0.4, "WebRTC cifrado de punto a punto · cambios agrupados cada 300 ms · reacciones en la pantalla en ~120 ms", size=13, color=MUTED, align=PP_ALIGN.CENTER)
nodo(3.2, 5.5, 6.9, 1.5, "Probado con 150 personas a la vez", "Todas entran, sus 150 votos se cuentan en 0,3 s y 450 reacciones llegan a la diapositiva en 0,17 s.", color=CARD2)

# ---------------------------------------------------------------- 12. privacidad

s = nueva()
cabecera(s, "Privacidad", "Las preguntas son de tu sala")
for i, (t, c) in enumerate([
    ("Sin cuentas ni analítica", "Ni el público ni el presentador se registran. Sin anuncios ni cookies."),
    ("Directo y cifrado", "Nada pasa por servidores del autor: solo se usa un servidor para encontrarse."),
    ("Nombre opcional", "Cada persona decide si pregunta con su nombre; nadie ve quién votó qué."),
    ("Solo en tu Chrome", "Las preguntas se guardan en tu equipo; las borras o exportas cuando quieras."),
]):
    tarjeta(s, 0.7 + i * 3.03, 2.25, 2.83, 3.6, t, c, numero=i + 1)

# ---------------------------------------------------------------- 13. estado y publicación

s = nueva()
cabecera(s, "Estado · versión 1.0.0", "Listo para publicar")
for i, (t, c) in enumerate([
    ("42 comprobaciones", "Extensión, panel, presentación con la capa y tres celulares reales: preguntas, votos, moderación, encuestas, reacciones y reconexión."),
    ("Paquete listo", "dist/preguntazo-1.0.0.zip para la Chrome Web Store, con permisos mínimos y acceso a sitios opcional."),
    ("Ficha completa", "Textos, 5 capturas, 2 mosaicos, política de privacidad y justificación de cada permiso."),
]):
    tarjeta(s, 0.7 + i * 4.05, 2.25, 3.85, 2.6, t, c, numero=i + 1)
texto(s, 0.7, 5.35, 11.9, 1.2, ["Siguiente paso: publicar la página del público en mjmc4498.github.io/preguntazo y subir el ZIP 1.0.0 en el panel del desarrollador."], size=16, color=MUTED, line=1.3)

# ---------------------------------------------------------------- 14. cierre

s = nueva()
foto(s, IMG / "logo-512.png", 0.7, 1.7, w=1.25)
texto(s, 0.7, 3.15, 8, 1.1, "¿Preguntas?", size=56, bold=True)
texto(s, 0.7, 4.3, 7.4, 1.4, ["Escanea el QR y haz tu pregunta: así de fácil será en tu próxima charla.", "mjmc4498.github.io/preguntazo"], size=20, color=MUTED, line=1.25)
texto(s, 0.7, 6.55, 7, 0.4, "Creado por Miguel Mogrovejo", size=13, color=MUTED)
foto(s, phones["phone-bienvenida"], 9.15, 0.55, h=6.4)

prs.save(OUT)
print("  ", OUT.relative_to(ROOT), f"({len(prs.slides)} diapositivas)")
